# -*- coding: utf-8 -*-
"""入库管线：各文件类型 → markdown → PageIndex 树 → DocStore。

模式移植自 kb-platform/ingest/pipelines.py（已验证），适配空间/库模型。
每个文档一个后台 job，进度事件写 events 表供 UI 轮询/SSE。
"""
import os
import threading
import traceback

from . import audio, config, db, kb, llm

JOBS = {}   # doc_id -> {"done": bool, "error": str|None}
ADHOC_INLINE_MAX = 24000   # 临时文件小于此字数则内联直答，否则建索引
INGEST_CONCURRENCY = 2     # 同时索引的文档数上限（批量上传防打爆模型端点）

_job_q = None
_workers_started = False


def kind_of(name):
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    return {"docx": "docx", "pptx": "pptx", "xlsx": "xlsx", "xls": "xls",
            "pdf": "pdf", "md": "md", "txt": "txt",
            "mp3": "audio", "wav": "audio", "m4a": "audio", "mp4": "audio",
            "aac": "audio", "flac": "audio", "ogg": "audio", "wma": "audio",
            "png": "image", "jpg": "image", "jpeg": "image",
            "webp": "image"}.get(ext, "unknown")


def docx_to_md(path):
    import docx
    d = docx.Document(path)
    out, tables = [], list(d.tables)
    ti = 0

    def para_iter(doc):
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        body = doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, doc)
            elif child.tag == qn("w:tbl"):
                yield Table(child, doc)

    for item in para_iter(d):
        from docx.table import Table
        if isinstance(item, Table):
            rows = [[c.text.strip().replace("\n", " ").replace("|", "｜")
                     for c in r.cells] for r in item.rows]
            if not rows:
                continue
            out.append("\n**[表格 " + str(ti + 1) + "]**\n")
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            for r in rows[1:]:
                out.append("| " + " | ".join(r) + " |")
            ti += 1
            continue
        t = item.text.strip()
        if not t:
            continue
        style = (item.style.name or "").lower()
        if style.startswith("heading"):
            lvl = "".join(c for c in style if c.isdigit()) or "1"
            out.append("#" * min(int(lvl) + 1, 6) + " " + t)
        elif style == "title":
            out.append("# " + t)
        else:
            out.append(t)
    return "\n\n".join(out)


def pptx_to_md(path):
    from pptx import Presentation
    pr = Presentation(path)
    out = []
    for i, slide in enumerate(pr.slides, 1):
        texts = []
        for shp in slide.shapes:
            try:
                if shp.has_text_frame:
                    for para in shp.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            texts.append(t)
            except Exception:
                pass
        out.append(f"## 幻灯片{i}\n\n" + "\n".join(texts))
    return "\n\n".join(out)


def xlsx_to_md(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## 工作表：{ws.title}")
        rows = [list(r) for r in ws.iter_rows(values_only=True)][:2000]
        rows = [r for r in rows if any(str(c).strip() for c in r if c)]
        if not rows:
            continue

        def cell(v):
            if v is None:
                return ""
            if isinstance(v, float) and v == int(v):
                return str(int(v))
            return str(v).strip().replace("|", "｜").replace("\n", " ")[:200]
        header = rows[0]
        ncol = len(header)
        out.append("| " + " | ".join(cell(c) for c in header) + " |")
        out.append("|" + "---|" * ncol)
        for r in rows[1:]:
            r = (list(r) + [""] * ncol)[:ncol]
            out.append("| " + " | ".join(cell(c) for c in r) + " |")
    wb.close()
    return "\n\n".join(out)


OCR_PROMPT = ("这是文档的一页扫描图。请尽可能完整地把所有文字转写为 markdown"
              "（表格用 | 表格语法，保留型号与数字精度），图形示意信息用一句话描述。")


def vision_ocr_image(png_bytes):
    """视觉 OCR：图片 → markdown。用当前通道（mimo 支持 vision 时）。"""
    import base64
    import io
    from PIL import Image
    im = Image.open(io.BytesIO(png_bytes)).convert("RGB")
    if min(im.size) < 400:
        k = 400 / min(im.size)
        im = im.resize((int(im.width * k), int(im.height * k)), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        png_bytes = buf.getvalue()
    b64 = base64.b64encode(png_bytes).decode()
    ref = llm.model_ref()
    import litellm
    r = litellm.completion(
        model=ref, timeout=300, drop_params=True,
        messages=[{"role": "user", "content": [
            {"type": "text", "text": OCR_PROMPT},
            {"type": "image_url",
             "image_url": {"url": "data:image/png;base64," + b64}}]}])
    return r.choices[0].message.content or ""


# ────────────────────────── markdown → PI 树 ──────────────────────────

def md_to_pi(md_text, name):
    """md 文本 → (tree, pages, description)。"""
    import hashlib
    work = os.path.join(config.DATA, "work")
    os.makedirs(work, exist_ok=True)
    tmp = os.path.join(work, "_" + hashlib.md5(
        name.encode()).hexdigest()[:8] + ".md")
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(md_text)
    try:
        tree = kb.md_to_tree_async(tmp, llm.model_ref())
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass
    pages = kb.md_pages(md_text)
    structure = kb.tree_from_md_structure(tree["structure"], len(pages))
    return structure, pages, tree.get("doc_description") or ""


def tag_suggest(name, md_text):
    """LLM 依据内容建议标签（入库时默认勾选，用户可改）。"""
    try:
        obj, _ = llm.complete_json(
            [{"role": "system", "content":
              "为文档建议 2-5 个中文标签（主题/系统名/文档类型，短词）。"
              '只输出 JSON：{"tags": ["..."]}'},
             {"role": "user", "content": f"文件名：{name}\n内容开头：\n"
              + md_text[:1500]}], max_tokens=200)
        return [str(t)[:24] for t in obj.get("tags", [])][:6]
    except Exception:
        return []


# ────────────────────────── job 执行 ──────────────────────────

def start_job(doc_row_id):
    """入队后台索引；全局最多 INGEST_CONCURRENCY 个文档并行。"""
    global _job_q, _workers_started
    import queue as _queue
    if _job_q is None:
        _job_q = _queue.Queue()
    if not _workers_started:
        for _ in range(INGEST_CONCURRENCY):
            threading.Thread(target=_worker_loop, daemon=True).start()
        _workers_started = True
    JOBS[doc_row_id] = {"done": False, "error": None}
    _job_q.put(doc_row_id)


def _worker_loop():
    while True:
        jid = _job_q.get()
        try:
            row = db.one("SELECT * FROM documents WHERE id=?", (jid,))
            if row and row["status"] == "pending":
                _run(row)
        except Exception:
            traceback.print_exc()
            try:
                db.x("UPDATE documents SET status='failed', error=?, updated_at=? "
                     "WHERE id=?", ("worker exception", db.now(), jid))
            except Exception:
                pass
        finally:
            _job_q.task_done()


def _progress(row, stage, detail):
    db.log_event(row["space_id"], row["id"], "progress",
                 {"stage": stage, "detail": detail})


def _run(row):
    jid = row["id"]
    try:
        db.x("UPDATE documents SET status='indexing', updated_at=? WHERE id=?",
             (db.now(), jid))
        _progress(row, "parse", "解析文件为 markdown")
        path = row["source_path"]
        kind = row["kind"]
        space_id = row["space_id"]
        meta = {"library_id": row["library_id"], "title": row["name"],
                "doc_row": jid}

        if kind == "pdf":
            doc_id, pages_n, desc = kb.submit_pdf(space_id, path)
            store = kb.get_store(space_id)
            m = store.get_meta(doc_id) or {}
            m.setdefault("metadata", {}).update(meta)
            with store.lock():
                store.save_document(doc_id, m, store.get_tree(doc_id) or [],
                                    store.get_pages(doc_id) or [])
            _update_ready(jid, doc_id, pages_n, desc)
            _progress(row, "done", f"PDF 索引完成（{pages_n} 页）")
            return

        if kind == "docx":
            md = docx_to_md(path)
        elif kind == "pptx":
            md = pptx_to_md(path)
        elif kind == "xlsx" or kind == "xls":
            md = xlsx_to_md(path)
        elif kind == "md":
            md = open(path, encoding="utf-8", errors="replace").read()
        elif kind == "txt":
            md = "# " + row["name"] + "\n\n```\n" + open(
                path, encoding="utf-8", errors="replace").read() + "\n```"
        elif kind == "image":
            md = f"# {row['name']}\n\n（图像识别结果）\n\n" + \
                vision_ocr_image(open(path, "rb").read())
        elif kind == "audio":
            _progress(row, "asr", "录音转写中（固定脚本管线）")
            md = audio.transcribe_to_minutes_md(path, row["name"],
                                                lambda d: _progress(row, "asr", d))
        else:
            raise RuntimeError(f"暂不支持的类型: {kind}")

        if not md.strip():
            raise RuntimeError("解析结果为空（可能是纯扫描件或空文件）")

        _progress(row, "tree", f"构建层级索引树（{len(md)} 字符）")
        tree, pages, desc = md_to_pi(md, row["name"])
        _progress(row, "seed", f"写入索引存储（{len(pages)} 页）")
        doc_id = kb.seed_doc(space_id, row["name"], desc, tree, pages, meta)

        # 标签建议（失败不阻塞入库）
        suggested = tag_suggest(row["name"], md)
        if suggested:
            _apply_suggested_tags(jid, space_id, suggested)

        # 摘要钩子（演进输入，失败不阻塞）
        try:
            _hook_evolution(row, desc, md)
        except Exception as e:
            print("evolution hook fail:", e)

        _update_ready(jid, doc_id, len(pages), desc)
        _progress(row, "done", f"入库完成（{len(pages)} 页）")
    except Exception as e:
        err = f"{type(e).__name__}: {e}"
        traceback.print_exc()
        db.x("UPDATE documents SET status='failed', error=?, updated_at=? "
             "WHERE id=?", (err[:400], db.now(), jid))
        _progress(row, "failed", err[:200])
        JOBS[jid] = {"done": True, "error": err}
        return
    JOBS[jid] = {"done": True, "error": None}


def _apply_suggested_tags(doc_id, space_id, tags):
    for name in tags:
        t = db.one("SELECT id FROM tags WHERE space_id=? AND name=?",
                   (space_id, name))
        tid = t["id"] if t else db.new_id("tg-")
        if not t:
            db.x("INSERT INTO tags(id,space_id,name,auto,created_at) "
                 "VALUES(?,?,?,1,?)", (tid, space_id, name, db.now()))
        db.x("INSERT OR IGNORE INTO doc_tags(document_id,tag_id) VALUES(?,?)",
             (doc_id, tid))


def _update_ready(jid, doc_id, pages, desc):
    db.x("UPDATE documents SET doc_id=?, pages=?, description=?, "
         "status='ready', updated_at=? WHERE id=?",
         (doc_id, pages, (desc or "")[:1000], db.now(), jid))


def _hook_evolution(row, desc, md):
    """入库演进钩子：会议纪要 → 抽取候选（决议/能力/术语）；更新库摘要。"""
    space_id = row["space_id"]
    db.x("INSERT OR REPLACE INTO profiles(id,space_id,name,content,ttl_days,"
         "updated_at) VALUES(?,?,?,?,90,?)",
         (db.new_id("pf-"), space_id, "doc:" + row["id"],
          (desc or md[:200])[:500], db.now()))
    # 库摘要增量重写
    lib = db.one("SELECT * FROM libraries WHERE id=?", (row["library_id"],))
    if lib:
        obj, _ = llm.complete_json(
            [{"role": "system", "content":
              '合并旧库摘要与新文档信息，输出更新后的库摘要（150字内，'
              '概括该库覆盖的主题与能力）。只输出 JSON：{"summary":"..."}'},
             {"role": "user", "content":
              f"库名：{lib['name']}\n旧摘要：{lib['summary'] or '（空）'}\n"
              f"新文档：{row['name']} — {desc or md[:300]}"}],
            max_tokens=300)
        db.x("UPDATE libraries SET summary=? WHERE id=?",
             ((obj.get("summary") or "")[:500], row["library_id"]))
    # 疑似会议纪要：抽取候选规则/能力
    low = row["name"].lower()
    if any(k in row["name"] + low for k in ("纪要", "会议", "meeting", "录音")):
        obj, _ = llm.complete_json(
            [{"role": "system", "content":
              '从会议纪要中抽取可沉淀为部门知识的事实/规则/系统能力声明。'
              '每条一句话陈述。没有则空数组。'
              '只输出 JSON：{"candidates":[{"kind":"rule|capability|glossary",'
              '"statement":"..."}]}，最多 6 条。'},
             {"role": "user", "content": md[:6000]}], max_tokens=600)
        for c in obj.get("candidates", [])[:6]:
            add_candidate(space_id, c.get("kind", "rule"),
                          c.get("statement", ""),
                          evidence=[{"doc": row["name"], "doc_id": row["id"]}],
                          source="ingest:" + row["id"])


def add_candidate(space_id, kind, statement, evidence=None, source="",
                  ex_bad="", ex_good=""):
    """候选写入 + 简易去重（同键合并 times_seen）。"""
    if not statement or len(statement) < 6:
        return None
    key = (kind + "|" + statement)[:120].lower()
    old = db.one("SELECT id FROM candidates WHERE space_id=? AND dedup_key=? "
                 "AND status='pending'", (space_id, key))
    if old:
        db.x("UPDATE candidates SET times_seen=times_seen+1, updated_at=? "
             "WHERE id=?", (db.now(), old["id"]))
        return old["id"]
    cid = db.new_id("cd-")
    import json
    db.x("INSERT INTO candidates(id,space_id,kind,statement,ex_bad,ex_good,"
         "evidence,source,status,dedup_key,created_at,updated_at) "
         "VALUES(?,?,?,?,?,?,?,?, 'pending',?,?,?)",
         (cid, space_id, kind, statement[:600], ex_bad[:300], ex_good[:300],
          json.dumps(evidence or [], ensure_ascii=False), source, key,
          db.now(), db.now()))
    return cid
