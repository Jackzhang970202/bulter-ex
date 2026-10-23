# -*- coding: utf-8 -*-
"""Extract text and images from Word/PPT in document order; images go through vision understanding."""
from . import vision


def _blips(el):
    from docx.oxml.ns import qn
    return el.findall('.//' + qn('a:blip'))


def _docx_image(doc, blip, doc_name, prev_text):
    from docx.oxml.ns import qn
    rid = blip.get(qn('r:embed'))
    part = doc.part.related_parts.get(rid) if rid else None
    if not part:
        return None
    try:
        rec = vision.identify_image(part.blob, doc_name, 0,
                                    prev_text=prev_text, hint="word")
    except Exception as e:
        rec = {"type": "other", "transcription": "", "status": "failed",
               "description": f"图片识别失败：{str(e)[:120]}", "sha": "", "img_path": ""}
    return rec


def docx_extract(path: str, doc_name: str, vision_on: bool, progress=None):
    """Return (markdown, images). Images are inserted in body XML order."""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    out, images, seq, prev_text = [], [], 0, ""

    def append_image(rec):
        nonlocal seq
        seq += 1
        md = vision.image_markdown(rec, seq=seq)
        out.append(md)
        images.append({
            "page": 0, "order_no": seq, "img_type": rec.get("type", ""),
            "transcription": rec.get("transcription", ""),
            "description": rec.get("description", ""),
            "status": rec.get("status", "ok"),
            "sha": rec.get("sha", ""), "img_path": rec.get("img_path", ""),
            "marker": md,
        })
        if progress:
            progress("ocr", f"图片理解 第{seq}张")

    def handle_paragraph(p):
        nonlocal prev_text
        t = p.text.strip()
        if t:
            style = (p.style.name or "").lower()
            if style.startswith("heading"):
                lvl = "".join(c for c in style if c.isdigit()) or "1"
                out.append("#" * min(int(lvl) + 1, 6) + " " + t)
            elif style == "title":
                out.append("# " + t)
            else:
                out.append(t)
            prev_text = t
        if vision_on:
            for blip in _blips(p._p):
                rec = _docx_image(d, blip, doc_name, prev_text)
                if rec:
                    append_image(rec)

    for child in d.element.body.iterchildren():
        if child.tag == qn("w:p"):
            handle_paragraph(Paragraph(child, d))
        elif child.tag == qn("w:tbl"):
            tb = Table(child, d)
            rows = [[c.text.strip().replace("\n", " ").replace("|", "｜")
                     for c in r.cells] for r in tb.rows]
            if rows:
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * len(rows[0]))
                for r in rows[1:]:
                    out.append("| " + " | ".join(r) + " |")
                if vision_on:
                    for row in tb.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                for blip in _blips(p._p):
                                    rec = _docx_image(d, blip, doc_name, prev_text)
                                    if rec:
                                        append_image(rec)
    return "\n\n".join(out), images


def pptx_extract(path: str, doc_name: str, vision_on: bool, progress=None):
    """Return (markdown, images). Shapes within a slide are processed in (top,left) order."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    out, images, seq, prev_text = [], [], 0, ""

    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## 幻灯片{i}")
        shapes = sorted(slide.shapes,
                        key=lambda s: (s.top or 0, s.left or 0))
        for shp in shapes:
            try:
                if shp.has_text_frame:
                    texts = []
                    for para in shp.text_frame.paragraphs:
                        t = "".join(run.text for run in para.runs).strip()
                        if t:
                            texts.append(t)
                    if texts:
                        joined = "\n".join(texts)
                        out.append(joined)
                        prev_text = joined
            except Exception:
                pass
            is_pic = shp.shape_type == MSO_SHAPE_TYPE.PICTURE or \
                hasattr(shp, "image")
            if is_pic and vision_on:
                try:
                    blob = shp.image.blob
                except Exception:
                    continue
                rec = vision.identify_image(blob, doc_name, i,
                                            prev_text=prev_text, hint="ppt")
                seq += 1
                md = vision.image_markdown(rec, page=i, seq=seq)
                out.append(md)
                images.append({
                    "page": i, "order_no": seq,
                    "img_type": rec.get("type", ""),
                    "transcription": rec.get("transcription", ""),
                    "description": rec.get("description", ""),
                    "status": rec.get("status", "ok"),
                    "sha": rec.get("sha", ""), "img_path": rec.get("img_path", ""),
                    "marker": md,
                })
                if progress:
                    progress("ocr", f"幻灯片{i} 图片理解 第{seq}张")
            elif shp.shape_type == MSO_SHAPE_TYPE.CHART or \
                    getattr(shp, "has_chart", False):
                out.append("（原生图表对象，未做视觉识别）")
    return "\n\n".join(out), images