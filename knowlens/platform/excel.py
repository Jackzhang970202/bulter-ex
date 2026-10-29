# -*- coding: utf-8 -*-
"""Excel natural-language query (问数).

The prompt lineage comes from the AskBI project (`askexcel_*` agents). Unlike
AskBI this module never calls `exec` in the API process: generated pandas code
is AST-checked and executed in a short-lived `python -I` subprocess with a
timeout, a minimal environment and no network-capable names.

Only Excel/CSV attachments are supported; database querying is intentionally
out of scope.
"""
import ast
import json
import os
import re
import shutil
import subprocess
import sys

import pandas as pd

from .. import config
from . import context as pf_context
from . import files as pf_files, llm_client, services

EXCEL_KINDS = {"xlsx", "xls", "csv"}
ROOT = os.path.join(config.DATA, "platform_excel")
MAX_ROWS = 5000
MAX_COLS = 200
MAX_CELLS = 200000
CODE_TIMEOUT = 30
MAX_CODE = 20000

CODE_SYSTEM = """你是 Excel 问数场景下的 Python 数据分析专家，擅长用 pandas 做筛选、聚合、统计与解释。
执行环境已预置 pd、json、datetime、re、math、statistics、FILE_LIST、FILE_METADATA、RESULT。
只输出 Python 代码本身，不要 markdown 代码块，不要解释。
必须把最终结果赋值给 RESULT，代码只依赖 FILE_LIST 中的文件。
禁止 import、禁止 open/exec/eval、禁止读写文件、禁止网络、禁止系统命令。
计数问题（多少/数量/几条/几人）用 len() 或 shape[0]，不要对 ID/序号求和。
结果是 DataFrame/Series 时先转成 list/dict 再赋给 RESULT。
列名必须来自 FILE_METADATA，不要臆造。"""

REPORT_SYSTEM = """你是业务数据分析师。结论必须基于真实结果数据；
不要输出代码、SQL、图表配置或 JSON；保留原始数值精度，不要擅自四舍五入；
用中文 Markdown 输出三节：## 1. 问题结果 / ## 2. 回答依据 / ## 3. 分析解读。"""

SYNTH_SYSTEM = """你是企业知识助手。你会同时拿到两类材料：
表格真实计算结果（数字）与知识库片段（制度、定义、口径）。
要求：
1. 数字只能来自表格结果，不得编造；说明用了哪些表/列与计算方式。
2. 制度、定义、口径来自知识库，必须标注文档名（及页码/小节）。
3. 严格区分来源：知识库里的内容绝不能写成“来自表格”，表格里的数字也不能写成“来自知识库”。
4. 两者冲突时明确指出冲突点，并说明以哪个为准、为什么。
5. 用中文 Markdown 输出：## 1. 结论 / ## 2. 数据依据 / ## 3. 制度依据 / ## 4. 建议（可选）。
6. 知识库没有覆盖的部分，明确写“知识库未覆盖”，不要臆测。"""

CHART_SYSTEM = """你是数据可视化助手。只输出一个 ECharts option JSON 对象，
不要 markdown、不要注释、不要函数或回调。数据必须内联在 series.data 中；
没有合适图表时输出 {"chart_needed": false, "reason": "..."}。"""


# ────────────────────────── profiling ──────────────────────────

def _numeric_ratio(row):
    vals = [v for v in row if v not in (None, "")]
    if not vals:
        return 0.0
    nums = 0
    for v in vals:
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            nums += 1
    return nums / len(vals)


def _clean_header(value, index):
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:60] or f"col{index + 1}"


def _dedupe(names):
    seen, out = {}, []
    for name in names:
        count = seen.get(name, 0) + 1
        seen[name] = count
        out.append(name if count == 1 else f"{name}_{count}")
    return out


def _sheet_grid(ws):
    """Read an openpyxl sheet into a 2-D list with merged cells expanded."""
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row[:MAX_COLS]))
        if len(rows) >= MAX_ROWS:
            break
    for merged in ws.merged_cells.ranges:
        value = None
        try:
            value = rows[merged.min_row - 1][merged.min_col - 1]
        except Exception:
            continue
        if value in (None, ""):
            continue
        for r in range(merged.min_row - 1, min(merged.max_row, len(rows))):
            for c in range(merged.min_col - 1, min(merged.max_col, MAX_COLS)):
                if rows[r][c] in (None, ""):
                    rows[r][c] = value
    return rows


def _grid_to_frame(grid):
    if not grid:
        return None
    header_idx = 0
    for i, row in enumerate(grid[:15]):
        nonempty = [v for v in row if v not in (None, "")]
        if len(nonempty) >= 2 and _numeric_ratio(row) < 0.5:
            header_idx = i
            break
    header = list(grid[header_idx])
    # A second header row is common in Chinese reports; join it when the first
    # header still has blanks and the next row is mostly textual.
    if header_idx + 1 < len(grid):
        nxt = grid[header_idx + 1]
        blanks = sum(1 for v in header if v in (None, ""))
        if blanks and nxt and _numeric_ratio(nxt) < 0.34:
            joined = []
            for i, (a, b) in enumerate(zip(header, nxt)):
                left = str(a).strip() if a not in (None, "") else ""
                right = str(b).strip() if b not in (None, "") else ""
                joined.append("_".join(x for x in (left, right) if x) or f"col{i + 1}")
            header = joined
            header_idx += 1
    names = _dedupe([_clean_header(v, i) for i, v in enumerate(header)])
    body = [row for row in grid[header_idx + 1:]
            if any(v not in (None, "") for v in row)]
    if not body:
        return None
    width = min(len(names), MAX_COLS)
    body = [row[:width] for row in body]
    df = pd.DataFrame(body, columns=names[:width])
    df = df.dropna(how="all").dropna(axis=1, how="all")
    for col in df.columns:
        df[col] = df[col].map(lambda v: v.strip() if isinstance(v, str) else v)
    return df if len(df.columns) else None


def profile(attachment_id, path):
    """Normalize every sheet into a clean xlsx and cache the profile."""
    folder = os.path.join(ROOT, attachment_id)
    cache = os.path.join(folder, "profile.json")
    if os.path.exists(cache):
        try:
            with open(cache, encoding="utf-8") as fh:
                return json.load(fh)
        except Exception:
            pass
    os.makedirs(folder, exist_ok=True)
    kind = (os.path.splitext(path)[1] or "").lower().lstrip(".")
    sheets = []
    if kind == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=False)
        for idx, ws in enumerate(wb.worksheets):
            df = _grid_to_frame(_sheet_grid(ws))
            if df is None or df.empty:
                continue
            sheets.append((ws.title or f"Sheet{idx + 1}", df))
    elif kind == "xls":
        raw_sheets = pd.read_excel(path, sheet_name=None, header=None,
                                   dtype=object)
        for idx, (name, raw) in enumerate(raw_sheets.items()):
            grid = raw.where(pd.notnull(raw), None).values.tolist()
            df = _grid_to_frame(grid)
            if df is None or df.empty:
                continue
            sheets.append((str(name or f"Sheet{idx + 1}"), df))
    else:
        csv_path = path
        try:
            df = pd.read_csv(csv_path, dtype=str)
        except UnicodeDecodeError:
            df = pd.read_csv(csv_path, dtype=str, encoding="gbk")
        df = df.dropna(how="all")
        sheets.append(("Sheet1", df))
    meta = []
    total_cells = 0
    for idx, (name, df) in enumerate(sheets):
        df = df.iloc[:MAX_ROWS, :MAX_COLS]
        total_cells += df.shape[0] * df.shape[1]
        if total_cells > MAX_CELLS:
            break
        safe = re.sub(r"[^0-9A-Za-z_一-龥-]+", "_", name)[:40] or f"sheet{idx+1}"
        out_path = os.path.join(folder, f"{idx + 1}_{safe}.xlsx")
        df.to_excel(out_path, index=False)
        meta.append({
            "sheet": name,
            "file": out_path,
            "columns": [str(c) for c in df.columns],
            "row_count": int(df.shape[0]),
            "sample_rows": json.loads(
                df.head(3).to_json(orient="records", force_ascii=False)),
        })
    data = {"attachment_id": attachment_id, "sheets": meta,
            "total_cells": total_cells,
            "truncated": total_cells >= MAX_CELLS}
    with open(cache, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False)
    return data


def profile_public(attachment_id, path):
    data = profile(attachment_id, path)
    return {"sheets": [{"name": s["sheet"], "columns": s["columns"],
                        "rows": s["row_count"],
                        "sample": s["sample_rows"]}
                       for s in data["sheets"]],
            "truncated": data.get("truncated")}


def suggestions(attachment_id, path):
    data = profile(attachment_id, path)
    out = []
    for sheet in data["sheets"][:2]:
        cols = sheet["columns"]
        for col in cols[:3]:
            out.append(f"按「{col}」统计数量")
            out.append(f"「{col}」的分布情况")
        if len(cols) >= 2:
            out.append(f"「{cols[0]}」和「{cols[1]}」有什么关系")
        out.append("这张表有多少行数据？")
    seen, result = set(), []
    for q in out:
        if q not in seen:
            seen.add(q)
            result.append(q)
    return result[:6]


# ────────────────────────── safe execution ──────────────────────────

BANNED_NAMES = {
    "open", "exec", "eval", "compile", "__import__", "getattr", "setattr",
    "delattr", "globals", "locals", "vars", "dir", "input", "breakpoint",
    "help", "os", "sys", "subprocess", "socket", "pathlib", "shutil",
    "importlib", "builtins", "pickle", "ctypes", "requests", "urllib",
}
BANNED_ATTRS = {
    "system", "popen", "spawn", "spawnl", "spawnv", "remove", "unlink",
    "rmdir", "mkdir", "makedirs", "rename", "replace_file", "write",
    "writelines", "to_excel", "to_csv", "to_pickle", "to_sql", "to_json",
    "savefig", "plot", "hist", "boxplot", "read_pickle", "read_sql",
    "read_html", "read_xml", "read_json", "read_table", "read_fwf",
    "read_clipboard", "read_hdf", "read_parquet", "read_feather",
    "read_orc", "read_stata", "read_sas", "read_spss", "read_gbq",
}
PATH_RE = re.compile(r"(https?://|ftp://|file://|^[A-Za-z]:[\\/]|^/|^\\)")
SUSPICIOUS_STR = re.compile(r"(^|[\\/])\.\.([\\/]|$)|__")


def validate_code(code):
    try:
        tree = ast.parse(code, mode="exec")
    except SyntaxError as exc:
        raise ValueError(f"语法错误：{exc}")
    for node in ast.walk(tree):
        if isinstance(node, (ast.Import, ast.ImportFrom)):
            raise ValueError("不允许 import")
        if isinstance(node, (ast.AsyncFunctionDef, ast.ClassDef,
                             ast.Global, ast.Nonlocal, ast.Delete,
                             ast.AsyncWith)):
            raise ValueError(f"不允许的语法：{type(node).__name__}")
        if isinstance(node, ast.arg) and node.arg in BANNED_NAMES:
            raise ValueError(f"不允许的参数名：{node.arg}")
        if isinstance(node, ast.Name) and node.id in BANNED_NAMES:
            raise ValueError(f"不允许的名称：{node.id}")
        if isinstance(node, ast.Attribute):
            if node.attr.startswith("_") or node.attr in BANNED_ATTRS:
                raise ValueError(f"不允许的属性：{node.attr}")
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            value = node.value.strip()
            if PATH_RE.search(value) or SUSPICIOUS_STR.search(value):
                raise ValueError("不允许的字符串（路径/URL/双下划线）")
    return tree


PRELUDE = """# -I ignores PYTHON* env vars, so pin UTF-8 explicitly for the child process.
import sys as _sys
for _stream in (_sys.stdout, _sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
import pandas as pd, json, datetime, re, math, statistics
FILE_LIST = {files}
FILE_METADATA = {metadata}
RESULT = None
{code}
def _safe(o):
    if hasattr(o, "to_dict"):
        try:
            return o.to_dict(orient="records")
        except Exception:
            try:
                return o.to_dict()
            except Exception:
                return str(o)
    return str(o)
print("__KP_RESULT__" + json.dumps(RESULT, ensure_ascii=False, default=_safe))
"""


def run_code(code, file_paths, metadata):
    tree = validate_code(code)
    folder = os.path.join(ROOT, "_run")
    os.makedirs(folder, exist_ok=True)
    script = os.path.join(folder, "job.py")
    payload = PRELUDE.format(
        files=repr(list(file_paths)),
        metadata=repr(json.dumps(metadata, ensure_ascii=False)),
        code=code[:MAX_CODE])
    with open(script, "w", encoding="utf-8") as fh:
        fh.write(payload)
    env = {"PATH": os.environ.get("PATH", ""),
           "SYSTEMROOT": os.environ.get("SYSTEMROOT", ""),
           "TEMP": os.environ.get("TEMP", folder),
           "TMP": os.environ.get("TMP", folder),
           "PYTHONIOENCODING": "utf-8",
           "PYTHONUTF8": "1"}
    try:
        proc = subprocess.run([sys.executable, "-I", "-X", "utf8", script],
                              cwd=folder,
                              env=env, capture_output=True, text=True,
                              encoding="utf-8", errors="replace",
                              timeout=CODE_TIMEOUT)
    except subprocess.TimeoutExpired:
        return {"ok": False, "error": f"执行超时（{CODE_TIMEOUT}s）"}
    out = proc.stdout or ""
    marker = "__KP_RESULT__"
    if marker in out:
        raw = out.split(marker, 1)[1].strip().splitlines()[0]
        try:
            value = json.loads(raw)
        except Exception:
            value = raw[:8000]
        return {"ok": True, "result": value, "stdout": out[:2000]}
    error = (proc.stderr or out or "没有返回结果").strip()
    return {"ok": False, "error": error[-2000:]}


# ────────────────────────── prompts ──────────────────────────

def _metadata_text(data, max_chars=6000):
    lines = []
    for sheet in data["sheets"]:
        lines.append(f"文件：{os.path.basename(sheet['file'])} | 工作表：{sheet['sheet']} | 行数：{sheet['row_count']}")
        lines.append("列名：" + "、".join(sheet["columns"]))
        lines.append("示例：" + json.dumps(sheet["sample_rows"][:2], ensure_ascii=False))
    text = "\n".join(lines)
    if len(text) > max_chars:
        text = text[:max_chars] + "\n…（元数据已截断）"
    return text


def _code_prompt(question, metadata_text, previous_error="", previous_code="",
                 history=None, kb_text=""):
    ctx = ""
    if history:
        ctx = "\n最近对话（用于理解追问）：\n" + "\n".join(
            f"问：{h.get('q')}\n答：{(h.get('a') or '')[:200]}"
            for h in history[-3:])
    kb = f"\n业务口径参考（来自知识库，仅用于确认列的业务含义）：\n{kb_text[:1200]}\n" \
        if kb_text else ""
    text = f"""用户问题：{question}
{ctx}

文件元数据：
{metadata_text}
{kb}

请只输出一段 Python 代码（不要 markdown 代码块、不要解释）：
1. 使用 pd.read_excel(FILE_LIST[i]) 读取；只能用 FILE_LIST 中的路径。
2. 必须使用元数据里的真实列名，不要臆造。
3. 计数问题用 len()/shape[0]，不要对 ID 或序号求和。
4. 数值列可能是字符串，需要 pd.to_numeric(errors="coerce")。
5. 结果赋值给 RESULT；DataFrame/Series 先转 list/dict。
6. 不要 import、不要读写文件、不要访问网络。
"""
    if previous_error:
        text += f"\n上一轮代码失败：{previous_error}\n请修复后输出完整代码。\n上一轮代码：\n{previous_code[:4000]}"
    return text


def _report_prompt(question, metadata_text, result, history):
    ctx = ""
    if history:
        ctx = "\n最近对话：\n" + "\n".join(
            f"问：{h.get('q')}\n答：{(h.get('a') or '')[:300]}" for h in history[-3:])
    return f"""用户问题：{question}
文件元数据：
{metadata_text}
真实执行结果：
{json.dumps(result, ensure_ascii=False, default=str)[:8000]}
{ctx}

请按以下结构输出中文 Markdown：
## 1. 问题结果
直接回答，给出关键数字。
## 2. 回答依据
说明用了哪些表/列、筛选条件与计算方式，必要时给出小表格。
## 3. 分析解读
两到三段，解释结果并指出需要注意的口径。"""


def _kb_answer(question, hits, kb_text, history=None):
    """Answer from knowledge only — used when the table cannot answer."""
    report = llm_client.complete([
        {"role": "system", "content": SYNTH_SYSTEM},
        {"role": "user", "content":
         f"用户问题：{question}\n\n表格无法给出结果（或问题与数据无关）。"
         f"请仅依据知识库片段回答，并标注出处。\n\n知识库片段：\n{kb_text}"},
    ], temperature=0.2, max_tokens=2000)
    return report.get("text") or ""


def ask(attachment_id, question, history=None, on_stage=None, kb_ids=None):
    row = pf_files.get(attachment_id)
    if not row:
        return {"ok": False, "error": "附件不存在"}
    path = row.get("path") or ""
    if not os.path.exists(path):
        return {"ok": False, "error": "附件文件不存在"}
    if (row.get("kind") or "") not in EXCEL_KINDS:
        return {"ok": False, "error": "只支持 Excel / CSV 附件"}

    def stage(name, detail=""):
        if on_stage:
            on_stage(name, detail)

    plan = pf_context.plan(question, kb_ids=kb_ids, excel_ids=[attachment_id])
    kb_text, kb_hits = "", []
    if plan["use_kb"]:
        stage("knowledge", plan["reason"])
        kb_hits = pf_context.retrieve(question, kb_ids=kb_ids)
        _, kb_text = pf_context.render(kb_hits, budget=4000)

    stage("profile", "解析工作表与列")
    data = profile(attachment_id, path)
    if not data["sheets"]:
        return {"ok": False, "error": "没有可读的工作表"}
    metadata_text = _metadata_text(data)
    files = [s["file"] for s in data["sheets"]]
    metadata = [{"sheet": s["sheet"], "file": os.path.basename(s["file"]),
                 "columns": s["columns"], "row_count": s["row_count"]}
                for s in data["sheets"]]

    code, error, rounds = "", "", 0
    result = None
    for rounds in range(1, 5):
        stage("code_generation", f"生成分析代码（第 {rounds} 轮）")
        out = llm_client.complete([
            {"role": "system", "content": CODE_SYSTEM},
            {"role": "user", "content": _code_prompt(
                question, metadata_text, error, code, history, kb_text)},
        ], temperature=0.1, max_tokens=3000)
        raw = re.sub(r"^```(?:python)?\s*", "", (out.get("text") or "").strip())
        raw = re.sub(r"\s*```$", "", raw)
        code = raw
        stage("code_execution", "执行代码")
        try:
            run = run_code(code, files, metadata)
        except Exception as exc:
            run = {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
        if run.get("ok"):
            result = run.get("result")
            break
        error = run.get("error") or "未知错误"
    if result is None and not error:
        return {"ok": False, "error": "未能生成可执行的分析代码"}
    if result is None:
        # Adaptive fallback: the table could not answer, knowledge might.
        if kb_text:
            stage("report_generation", "表格无法回答，改用知识库")
            return {"ok": True, "mode": "kb",
                    "summary": _kb_answer(question, kb_hits, kb_text,
                                          history or []),
                    "code": code, "result": None, "chart": None,
                    "rounds": rounds, "error": error,
                    "citations": pf_context.citations(kb_hits),
                    "sheets": [{"name": s["sheet"], "rows": s["row_count"],
                                "columns": s["columns"]}
                               for s in data["sheets"]],
                    "trace": {"metadata": metadata, "rounds": rounds,
                              "plan": plan}}
        return {"ok": False, "error": error, "code": code, "rounds": rounds}

    stage("report_generation", "生成结论与依据")
    report = llm_client.complete([
        {"role": "system", "content": REPORT_SYSTEM},
        {"role": "user", "content": _report_prompt(
            question, metadata_text, result, history or [])},
    ], temperature=0.2, max_tokens=2500)
    summary = report.get("text") or ""
    if kb_text:
        stage("knowledge_merge", "合并表格结果与知识库依据")
        merged = llm_client.complete([
            {"role": "system", "content": SYNTH_SYSTEM},
            {"role": "user", "content":
             f"用户问题：{question}\n\n表格真实结果（JSON）：\n"
             f"{json.dumps(result, ensure_ascii=False, default=str)[:6000]}\n\n"
             f"数据口径与计算方式初稿：\n{summary[:2500]}\n\n"
             f"知识库片段（含文档名与页码）：\n{kb_text or '（无）'}"},
        ], temperature=0.2, max_tokens=2500)
        if merged.get("text"):
            summary = merged["text"]

    chart = None
    try:
        stage("chart_generation", "判断是否需要图表")
        chart_out = llm_client.complete([
            {"role": "system", "content": CHART_SYSTEM},
            {"role": "user", "content":
             f"用户问题：{question}\n结果数据："
             f"{json.dumps(result, ensure_ascii=False, default=str)[:6000]}\n"
             "如果需要图表，给出 ECharts option；否则返回 chart_needed=false。"},
        ], temperature=0.1, max_tokens=1200)
        obj = llm_client.parse_json(chart_out.get("text") or "")
        if isinstance(obj, dict) and obj.get("series") and obj.get("xAxis"):
            chart = obj
    except Exception:
        chart = None

    return {"ok": True, "mode": plan["mode"], "summary": summary,
            "code": code, "result": result, "chart": chart, "rounds": rounds,
            "citations": pf_context.citations(kb_hits),
            "sheets": [{"name": s["sheet"], "rows": s["row_count"],
                        "columns": s["columns"]} for s in data["sheets"]],
            "trace": {"metadata": metadata, "rounds": rounds, "plan": plan}}